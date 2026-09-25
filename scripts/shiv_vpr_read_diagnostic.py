#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rhesus SHIV — vpr-region Read Pull + Soft-clip Diagnostic (Read-Only)
====================================================================
The question this answers, and ONLY this question:

  For reads that map to the vpr barcode window, do they reach the cassette
  insertion point and then soft-clip the 34bp barcode that is missing from the
  reference? And do those clipped/interior bases actually carry the barcode
  flanks and a canonicalizable core we can match to the Binhua ground truth?

If yes, reference surgery (adding the cassette/flank at the insertion locus so
reads align through it) is the right move and we build it next.
If coverage over the window is ~0, the reads are not there and surgery rescues
nothing -> the honest path is targeted VpxF1 cDNA enrichment (new sequencing).
If reads read straight through with no clip and no flank, the anchor model is
wrong and we recompute it from the reads.

This script does NOT modify any reference, BAM, or object. It pulls a ~500bp
window from indexed BAMs (instant, index-seek), profiles it, and writes CSVs +
two figures + a verdict.

Sources pulled per library:
  1. STARsolo SHIV-only recovery BAM  (relaxed filters; the recovery ceiling)
  2. CellRanger multi per-sample BAM  (combined ref, strict filters; cell-linked)
  3. CellRanger multi unassigned BAM  (reads not assigned to a called cell)

Conventions: Spyder cells (# %%), config in Cell 1, warn-and-continue, read-only.
Runs headless under SLURM (HEADLESS=True) or interactively in Spyder (False).

Author: Jake Lehle / Kaushal Lab
Date: July 2026
"""

# %% ============================================================
# CELL 1 — CONFIG AND DIALS
# ============================================================
import os
import re
import subprocess
import shutil
from collections import Counter, defaultdict

import numpy as np
import pandas as pd

import matplotlib
HEADLESS = True                      # False in Spyder for interactive plt.show()
if HEADLESS:
    matplotlib.use('Agg')
import matplotlib.pyplot as plt

# --- paths -----------------------------------------------------------------
WORKING_DIR   = '/master/jlehle/WORKING/SC/fastq/Rhesus_SHIV'
STARSOLO_DIR  = os.path.join(WORKING_DIR, 'shiv_starsolo')
SOLO_DIR      = os.path.join(STARSOLO_DIR, 'solo')
REF_FA        = os.path.join(STARSOLO_DIR, 'ref', 'SHIVAD8EO.fa')
CELLRANGER_DIR = os.path.join(WORKING_DIR, 'Analysis')
BINHUA_XLSX   = os.path.join(WORKING_DIR, 'Binhua_FINAL_Barcode_4726.xlsx')
OUT_DIR       = os.path.join(WORKING_DIR, 'annotation_output', 'vpr_read_diagnostic')

# --- library map (authoritative) ------------------------------------------
LIB_NAME = {
    'A': 'SHIV_Pre_PBMC',             'B': 'SHIV_Pre_LN',
    'C': 'SHIV_Pre_PBMC_IndianMixed', 'D': 'SHIV_21DPI_Pre_LN_Mixed',
    'E': 'SHIV_21DPI_PBMC',           'F': 'SHIV_21DPI_NonInf_LN_Mixed',
    'G': 'SHIV_Necropsy_PBMC',        'H': 'SHIV_Necropsy_LN',
}
LIBS      = ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H']
INFECTED  = {'D', 'E', 'F', 'H'}
NEGATIVE  = {'A', 'B', 'C', 'G'}

# --- vpr window (0-based, from the prior Module 1 primer hits) --------------
# VpxF1 fwd at 5790, vpr RT primer rev at 6097 -> insert bracketed ~5790..6118.
# Pad generously; the exact insertion locus is pinned empirically in Cell 2.
VPXF1  = 'CTAGGGGAAGGACATGGGGCAGG'
VPR_RT = 'CAGGTTGGCCGATTCTGGAGT'
WIN_START0 = 5700          # 0-based inclusive
WIN_END0   = 6210          # 0-based exclusive
# samtools regions are 1-based inclusive:
WIN_START1 = WIN_START0 + 1
WIN_END1   = WIN_END0

# --- barcode geometry (Khanal et al.; confirmed vs Binhua) -----------------
BC_LEN        = 34
ANCHOR_STR    = 'ACGCGC.{22}GCGCGT'
ANCHOR_RE     = re.compile(ANCHOR_STR)
MLUI          = 'ACGCGT'            # MluI recognition; likely the insertion locus
INNER_ORIENT1 = 'AGCGTA'            # core taken as-is
INNER_ORIENT2 = 'TAGCTG'            # core reverse-complemented to canonicalize
FLANK_CONST_LEN = 12                # constant 12bp ends of each 34bp barcode

# --- CellRanger BAM options ------------------------------------------------
INCLUDE_CELLRANGER_SAMPLE = True    # per_sample_outs/<s>/count/sample_alignments.bam
INCLUDE_CELLRANGER_UNASSIGNED = True  # outs/unassigned_alignments.bam
SHIV_LEN_EXPECTED = 10262           # from prior run; verify per BAM before trusting coords

# --- samtools (never the broken /usr/local/bin v0.1.18) --------------------
SAMTOOLS = os.environ.get('SAMTOOLS') or shutil.which('samtools')

# --- figure style (Kaushal conventions) ------------------------------------
FONT_MIN, FONT_MAX = 28, 34
HEX_INF = '#C0392B'                 # infected
HEX_NEG = '#2C7FB8'                 # negative
HEX_INS = '#27AE60'                 # insertion-point marker
DPI = 300

os.makedirs(OUT_DIR, exist_ok=True)
pd.set_option('display.width', 180)
pd.set_option('display.max_columns', 60)

_COMP = str.maketrans('ACGTNacgtn', 'TGCANtgcan')
CIGAR_RE = re.compile(r'(\d+)([MIDNSHP=X])')


def revcomp(s):
    return s.translate(_COMP)[::-1]


def canon_core(bc34):
    """Forward-strand 34bp anchor match -> orientation-canonical 10bp core."""
    if len(bc34) != BC_LEN:
        return None
    inner = bc34[6:12]
    core = bc34[12:22]
    if inner == INNER_ORIENT1:
        return core
    if inner == INNER_ORIENT2:
        return revcomp(core)
    return None


def parse_cigar(cig):
    if cig == '*' or not cig:
        return []
    return [(int(n), op) for n, op in CIGAR_RE.findall(cig)]


def ref_span(ops):
    return sum(n for n, op in ops if op in 'MDN=X')


def lead_softclip(ops):
    return ops[0][0] if ops and ops[0][1] == 'S' else 0


def trail_softclip(ops):
    return ops[-1][0] if ops and ops[-1][1] == 'S' else 0


def max_insertion(ops):
    ins = [n for n, op in ops if op == 'I']
    return max(ins) if ins else 0


def hr(title):
    print('\n' + '=' * 84)
    print(title)
    print('=' * 84)


def save_csv(df, name, index=True):
    p = os.path.join(OUT_DIR, name)
    df.to_csv(p, index=index)
    print(f'    saved: {name}')


print('vpr read diagnostic — config')
print(f'  samtools : {SAMTOOLS}')
if SAMTOOLS:
    try:
        v = subprocess.run([SAMTOOLS, '--version'], capture_output=True, text=True)
        first = v.stdout.splitlines()[0] if v.stdout else '?'
        print(f'             {first}')
        if '0.1.18' in first or (SAMTOOLS or '').startswith('/usr/local/bin'):
            print('  *** WARNING: this looks like the broken system samtools. '
                  'Activate sc_pre and export CONDA_PREFIX/bin to PATH.')
    except Exception as e:
        print(f'  WARN: could not query samtools version ({e})')
else:
    print('  *** samtools not found on PATH; nothing will run. Activate sc_pre.')
print(f'  window   : SHIV:{WIN_START1}-{WIN_END1} (1-based)   out: {OUT_DIR}')


# %% ============================================================
# CELL 2 — MODULE 1 (reference window + signpost derivation)
# ============================================================
hr('CELL 2 — reference window content, insertion locus, and flank signposts')

ref_seq = ''
ref_name = None
if not os.path.exists(REF_FA):
    print(f'  WARN: reference not found at {REF_FA}; downstream coords unverified.')
else:
    chunks = []
    with open(REF_FA) as fh:
        for line in fh:
            if line.startswith('>'):
                ref_name = line[1:].strip().split()[0]
            else:
                chunks.append(line.strip().upper())
    ref_seq = ''.join(chunks)
    print(f'  contig {ref_name}: {len(ref_seq):,} bp')
    if len(ref_seq) != SHIV_LEN_EXPECTED:
        print(f'  *** WARNING: contig length {len(ref_seq)} != expected '
              f'{SHIV_LEN_EXPECTED}. Coordinates below may be off.')

    win = ref_seq[WIN_START0:WIN_END0]
    print(f'\n  reference window {WIN_START0}..{WIN_END0} (0-based), {len(win)} bp:')
    for i in range(0, len(win), 60):
        print(f'    {WIN_START0 + i:>6}  {win[i:i + 60]}')

    def locate(label, probe):
        fwd = [m.start() for m in re.finditer(re.escape(probe), ref_seq)]
        rev = [m.start() for m in re.finditer(re.escape(revcomp(probe)), ref_seq)]
        print(f'  {label:<18} fwd {fwd if fwd else "-"}   rev {rev if rev else "-"}')
        return fwd, rev

    print('\n  probe / motif positions (0-based):')
    locate('VpxF1', VPXF1)
    locate('vpr RT primer', VPR_RT)
    anch = [m.start() for m in ANCHOR_RE.finditer(ref_seq)]
    print(f'  {"anchor 34bp":<18} {anch if anch else "- (absent from reference, as expected)"}')
    mlui_fwd = [m.start() for m in re.finditer(MLUI, ref_seq) if WIN_START0 <= m.start() < WIN_END0]
    mlui_rev = [m.start() for m in re.finditer(revcomp(MLUI), ref_seq) if WIN_START0 <= m.start() < WIN_END0]
    print(f'  {"MluI in window":<18} fwd {mlui_fwd if mlui_fwd else "-"}   '
          f'rev {mlui_rev if mlui_rev else "-"}  (candidate insertion locus)')

    # Reference-side flanks around the candidate insertion locus (the "5'UTR"
    # flank Jake wants pinned). If an MluI site is present, use it; else fall
    # back to the midpoint of the primer bracket.
    if mlui_fwd:
        ins0 = mlui_fwd[0]
        basis = 'MluI site'
    else:
        vpxf = re.search(re.escape(VPXF1), ref_seq)
        vprr = re.search(re.escape(revcomp(VPR_RT)), ref_seq)
        ins0 = (vpxf.end() + vprr.start()) // 2 if (vpxf and vprr) else (WIN_START0 + WIN_END0) // 2
        basis = 'primer-bracket midpoint (no MluI found)'
    ref_flank5 = ref_seq[max(0, ins0 - 24):ins0]
    ref_flank3 = ref_seq[ins0:ins0 + 24]
    print(f'\n  candidate insertion locus (0-based): {ins0}   [{basis}]')
    print(f'    reference 5-prime flank (24bp): {ref_flank5}')
    print(f'    reference 3-prime flank (24bp): {ref_flank3}')


# %% ============================================================
# CELL 3 — Binhua ground truth + cassette flank signposts
# ============================================================
hr('CELL 3 — Binhua ground truth + cassette-side constant flanks')

gt_exact = defaultdict(set)     # animal -> {34bp, both orientations}
gt_core  = defaultdict(set)     # animal -> {canonical 10bp core}
prefix_counts = Counter()       # constant 12bp 5' ends of the 34bp barcode
suffix_counts = Counter()       # constant 12bp 3' ends
have_gt = False

LIB_ANIMALS = {
    'A': ['39272', '40702', '40707', '41861', '41862'],
    'B': ['39272', '40702', '40707', '41861', '41862'],
    'C': ['34315', '41903'],
    'D': ['39272', '40702', '40707', '41861', '41862'],
    'E': ['39272', '40702', '40707', '41861', '41862'],
    'F': ['34315', '41903'],
    'G': ['40702', '40707', '41861', '41862'],
    'H': ['40702', '40707', '41861', '41862'],
}

if not os.path.exists(BINHUA_XLSX):
    print(f'  WARN: Binhua file not found at {BINHUA_XLSX}; core cross-ref disabled.')
else:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(BINHUA_XLSX, read_only=True, data_only=True)
        for sh in wb.sheetnames:
            animal = sh.split()[0]
            ws = wb[sh]
            n = 0
            for row in ws.iter_rows(values_only=True):
                if row and row[0] is not None and str(row[0]).strip().isdigit():
                    bc = row[2]
                    if isinstance(bc, str):
                        bc = bc.strip().upper()
                        if len(bc) == BC_LEN and set(bc) <= set('ACGTN'):
                            gt_exact[animal].add(bc)
                            gt_exact[animal].add(revcomp(bc))
                            c = canon_core(bc) or canon_core(revcomp(bc))
                            if c:
                                gt_core[animal].add(c)
                            prefix_counts[bc[:FLANK_CONST_LEN]] += 1
                            suffix_counts[bc[-FLANK_CONST_LEN:]] += 1
                            n += 1
            print(f'  {animal}: {n} barcodes ({len(gt_core[animal])} canonical cores)')
        wb.close()
        have_gt = True
    except Exception as e:
        print(f'  WARN: could not parse Binhua ({e}); core cross-ref disabled.')

gt_core_all = set().union(*gt_core.values()) if gt_core else set()
if have_gt:
    print(f'\n  pooled: {len(gt_core_all)} unique canonical cores')
    print('  cassette-side constant 5-prime flanks (12bp) by frequency:')
    for seq, n in prefix_counts.most_common(6):
        print(f'    {seq}  x{n}')
    print('  cassette-side constant 3-prime flanks (12bp) by frequency:')
    for seq, n in suffix_counts.most_common(6):
        print(f'    {seq}  x{n}')

# The sequences we scan reads for. Anchor first (definitive); constant flanks
# as a softer signpost that survives when only one end of the cassette is in-read.
FLANK_PROBES = set()
for seq, _ in prefix_counts.most_common(4):
    FLANK_PROBES.add(seq); FLANK_PROBES.add(revcomp(seq))
for seq, _ in suffix_counts.most_common(4):
    FLANK_PROBES.add(seq); FLANK_PROBES.add(revcomp(seq))
FLANK_PROBES = sorted(FLANK_PROBES)
print(f'  {len(FLANK_PROBES)} flank probe sequences prepared (both orientations)')


# %% ============================================================
# CELL 4 — BAM manifest (STARsolo + CellRanger), with contig/coord verification
# ============================================================
hr('CELL 4 — BAM manifest + SHIV contig detection')


def detect_shiv_contig(bam):
    """Return (name, length) of the SHIV contig in this BAM, or (None, None)."""
    try:
        hd = subprocess.run([SAMTOOLS, 'view', '-H', bam],
                            capture_output=True, text=True).stdout
    except Exception as e:
        print(f'      WARN header read failed on {os.path.basename(bam)} ({e})')
        return None, None
    best = None
    for line in hd.splitlines():
        if not line.startswith('@SQ'):
            continue
        name = length = None
        for tok in line.split('\t'):
            if tok.startswith('SN:'):
                name = tok[3:]
            elif tok.startswith('LN:'):
                length = int(tok[3:])
        if name is None:
            continue
        if ('SHIV' in name.upper()) or ('AD8' in name.upper()):
            return name, length          # name match is decisive
        if length and 9000 <= length <= 12000:
            best = (name, length)         # length fallback
    return best if best else (None, None)


manifest = []   # dicts: lib, group, source, bam, contig, contig_len, coord_ok

# 1. STARsolo BAMs
for L in LIBS:
    bam = os.path.join(SOLO_DIR, f'{L}_{LIB_NAME[L]}_Aligned.sortedByCoord.out.bam')
    if not os.path.exists(bam):
        print(f'  {L} starsolo: MISSING {os.path.basename(bam)}')
        continue
    name, ln = detect_shiv_contig(bam)
    manifest.append({'lib': L, 'group': 'INF' if L in INFECTED else 'neg',
                     'source': 'starsolo', 'bam': bam, 'contig': name,
                     'contig_len': ln, 'coord_ok': (ln == SHIV_LEN_EXPECTED)})

# 2. CellRanger per-sample + unassigned BAMs
for L in LIBS:
    sample = LIB_NAME[L]
    base = os.path.join(CELLRANGER_DIR, sample, 'outs')
    candidates = []
    if INCLUDE_CELLRANGER_SAMPLE:
        p = os.path.join(base, 'per_sample_outs', sample, 'count', 'sample_alignments.bam')
        candidates.append(('cr_sample', p))
        # fallback glob if the standard path is not present
        if not os.path.exists(p):
            import glob
            alt = glob.glob(os.path.join(base, 'per_sample_outs', sample, '**', '*.bam'),
                            recursive=True)
            for a in alt:
                candidates.append(('cr_sample', a))
    if INCLUDE_CELLRANGER_UNASSIGNED:
        candidates.append(('cr_unassigned', os.path.join(base, 'unassigned_alignments.bam')))
    seen = set()
    for src, p in candidates:
        if p in seen or not os.path.exists(p):
            continue
        seen.add(p)
        name, ln = detect_shiv_contig(p)
        manifest.append({'lib': L, 'group': 'INF' if L in INFECTED else 'neg',
                         'source': src, 'bam': p, 'contig': name,
                         'contig_len': ln, 'coord_ok': (ln == SHIV_LEN_EXPECTED)})

man_df = pd.DataFrame(manifest)
if len(man_df):
    print(man_df[['lib', 'group', 'source', 'contig', 'contig_len', 'coord_ok']].to_string(index=False))
    bad = man_df[~man_df['coord_ok']]
    if len(bad):
        print(f'\n  *** {len(bad)} BAM(s) have a SHIV contig length != {SHIV_LEN_EXPECTED}.')
        print('      Their window coordinates are NOT trusted and are skipped in Cell 5.')
    save_csv(man_df, 'vpr_bam_manifest.csv', index=False)
else:
    print('  No BAMs found. Check paths in Cell 1.')


# %% ============================================================
# CELL 5 — coordinate pull + per-read features
# ============================================================
hr('CELL 5 — pull vpr-window reads; record CIGAR, soft-clips, anchors, cores')


def scan_bam(row):
    """Pull reads over the SHIV window from one BAM; return (reads_list, depth_array)."""
    bam, contig = row['bam'], row['contig']
    reads, depth = [], np.zeros(WIN_END0 - WIN_START0, dtype=int)
    if contig is None or not row['coord_ok']:
        return reads, depth
    region = f'{contig}:{WIN_START1}-{WIN_END1}'

    # per-base depth over the window (aligned bases only; soft-clips excluded)
    try:
        dp = subprocess.run([SAMTOOLS, 'depth', '-a', '-r', region, bam],
                            capture_output=True, text=True).stdout
        for line in dp.splitlines():
            f = line.split('\t')
            if len(f) >= 3:
                pos1 = int(f[1]); d = int(f[2])
                idx = (pos1 - 1) - WIN_START0
                if 0 <= idx < depth.size:
                    depth[idx] = d
    except Exception as e:
        print(f'      WARN depth failed on {os.path.basename(bam)} ({e})')

    # per-read features
    try:
        p = subprocess.Popen([SAMTOOLS, 'view', bam, region],
                             stdout=subprocess.PIPE, text=True, bufsize=1 << 20)
        for line in p.stdout:
            f = line.rstrip('\n').split('\t')
            if len(f) < 11:
                continue
            flag = int(f[1]); pos1 = int(f[3]); mapq = int(f[4])
            cig = f[5]; seq = f[9]
            ops = parse_cigar(cig)
            rspan = ref_span(ops)
            ref_end1 = pos1 + rspan - 1
            lc, tc = lead_softclip(ops), trail_softclip(ops)
            maxins = max_insertion(ops)

            cb = ub = None
            for tag in f[11:]:
                if tag.startswith('CB:Z:'):
                    cb = tag[5:]
                elif tag.startswith('UB:Z:'):
                    ub = tag[5:]

            # scan the full read (both orientations) for the definitive anchor
            core = None
            anchor_hit = False
            for s in (seq, revcomp(seq)):
                m = ANCHOR_RE.search(s)
                if m:
                    anchor_hit = True
                    c = canon_core(m.group(0))
                    if c:
                        core = c
                        break
            # softer flank signpost
            flank_hit = any(fp in seq for fp in FLANK_PROBES)

            gt_hit = bool(core and core in gt_core_all)

            # clip boundary in reference coords: where does the clip start/end
            clip5_boundary = pos1 - 1 if lc else None       # ref coord just before read start
            clip3_boundary = ref_end1 + 1 if tc else None   # ref coord just after read end

            reads.append({
                'lib': row['lib'], 'group': row['group'], 'source': row['source'],
                'pos1': pos1, 'ref_end1': ref_end1, 'mapq': mapq, 'cigar': cig,
                'lead_sc': lc, 'trail_sc': tc, 'max_ins': maxins,
                'anchor': int(anchor_hit), 'flank': int(flank_hit),
                'core': core or '', 'gt_match': int(gt_hit),
                'clip5_boundary': clip5_boundary, 'clip3_boundary': clip3_boundary,
                'cb': cb or '', 'ub': ub or '',
                'lead_sc_seq': seq[:lc] if lc else '',
                'trail_sc_seq': seq[-tc:] if tc else '',
            })
        p.wait()
    except Exception as e:
        print(f'      WARN view failed on {os.path.basename(bam)} ({e})')
    return reads, depth


all_reads = []
depth_by = defaultdict(lambda: np.zeros(WIN_END0 - WIN_START0, dtype=int))  # (source,group)->depth

for _, row in (man_df.iterrows() if len(man_df) else []):
    reads, depth = scan_bam(row)
    all_reads.extend(reads)
    depth_by[(row['source'], row['group'])] += depth
    print(f'  {row["lib"]} [{row["group"]}] {row["source"]:<14}: '
          f'{len(reads):5d} window reads, '
          f'{sum(r["anchor"] for r in reads):3d} anchor, '
          f'{sum(r["flank"] for r in reads):3d} flank, '
          f'{sum(bool(r["core"]) for r in reads):3d} core, '
          f'{sum(r["gt_match"] for r in reads):3d} gt-match, '
          f'max depth {int(depth.max())}')

reads_df = pd.DataFrame(all_reads)
if len(reads_df):
    save_csv(reads_df, 'vpr_window_reads.csv', index=False)
else:
    print('  No reads pulled from the window in any BAM.')


# %% ============================================================
# CELL 6 — profiles, figures, verdict
# ============================================================
hr('CELL 6 — coverage profile, clip-boundary histogram, verdict')

xcoords = np.arange(WIN_START0, WIN_END0)

# ---- Figure 1: coverage across the window (infected vs negative, per source) ----
try:
    fig, ax = plt.subplots(figsize=(16, 8))
    styles = {'starsolo': '-', 'cr_sample': '--', 'cr_unassigned': ':'}
    for (source, group), dep in sorted(depth_by.items()):
        if dep.max() == 0:
            continue
        color = HEX_INF if group == 'INF' else HEX_NEG
        ax.plot(xcoords, dep, styles.get(source, '-'), color=color, linewidth=2,
                label=f'{source} [{group}]')
    if 'ins0' in dir():
        ax.axvline(ins0, color=HEX_INS, linewidth=3, linestyle='-', alpha=0.7,
                   label='insertion locus')
    ax.set_xlabel('SHIVAD8EO position (0-based)', fontsize=FONT_MAX)
    ax.set_ylabel('aligned depth', fontsize=FONT_MAX)
    ax.set_title('vpr-window coverage', fontsize=FONT_MAX)
    ax.tick_params(labelsize=FONT_MIN)
    lg = ax.legend(fontsize=FONT_MIN - 8, loc='upper right')
    fig.tight_layout()
    for ext in ('pdf', 'png'):
        fig.savefig(os.path.join(OUT_DIR, f'vpr_coverage.{ext}'), dpi=DPI, bbox_inches='tight')
    print('    saved: vpr_coverage.pdf / .png')
    if not HEADLESS:
        plt.show()
    plt.close(fig)
except Exception as e:
    print(f'  WARN: coverage figure failed ({e})')

# ---- Figure 2: clip-boundary histogram (where clips start, infected vs neg) ----
try:
    fig, ax = plt.subplots(figsize=(16, 8))
    if len(reads_df):
        for group, color in (('INF', HEX_INF), ('neg', HEX_NEG)):
            sub = reads_df[reads_df['group'] == group]
            b = pd.concat([sub['clip5_boundary'].dropna(),
                           sub['clip3_boundary'].dropna()])
            b = b[(b >= WIN_START0) & (b < WIN_END0)]
            if len(b):
                ax.hist(b, bins=np.arange(WIN_START0, WIN_END0 + 1, 4),
                        color=color, alpha=0.55, label=f'clip boundary [{group}]')
    if 'ins0' in dir():
        ax.axvline(ins0, color=HEX_INS, linewidth=3, alpha=0.7, label='insertion locus')
    ax.set_xlabel('SHIVAD8EO position (0-based)', fontsize=FONT_MAX)
    ax.set_ylabel('soft-clipped reads', fontsize=FONT_MAX)
    ax.set_title('where vpr-window reads soft-clip', fontsize=FONT_MAX)
    ax.tick_params(labelsize=FONT_MIN)
    ax.legend(fontsize=FONT_MIN - 8, loc='upper right')
    fig.tight_layout()
    for ext in ('pdf', 'png'):
        fig.savefig(os.path.join(OUT_DIR, f'vpr_clip_boundary.{ext}'), dpi=DPI, bbox_inches='tight')
    print('    saved: vpr_clip_boundary.pdf / .png')
    if not HEADLESS:
        plt.show()
    plt.close(fig)
except Exception as e:
    print(f'  WARN: clip-boundary figure failed ({e})')

# ---- numeric summary per source x group ----
if len(reads_df):
    g = (reads_df.groupby(['source', 'group'])
         .agg(reads=('pos1', 'size'),
              anchor=('anchor', 'sum'),
              flank=('flank', 'sum'),
              cores=('core', lambda s: (s != '').sum()),
              gt=('gt_match', 'sum'),
              clipped=('trail_sc', lambda s: (s > 0).sum()),
              max_depth=('pos1', 'size'))  # placeholder, real depth below
         )
    print('\n  per source x group:')
    print(g.drop(columns='max_depth').to_string())
    save_csv(g.drop(columns='max_depth'), 'vpr_window_summary.csv')

# ---- verdict ----
lines = []


def say(s=''):
    print(s)
    lines.append(s)


max_depth_overall = max((d.max() for d in depth_by.values()), default=0)
inf_anchor = int(reads_df[reads_df['group'] == 'INF']['anchor'].sum()) if len(reads_df) else 0
neg_anchor = int(reads_df[reads_df['group'] == 'neg']['anchor'].sum()) if len(reads_df) else 0
inf_gt = int(reads_df[reads_df['group'] == 'INF']['gt_match'].sum()) if len(reads_df) else 0
inf_clip = int(reads_df[(reads_df['group'] == 'INF') & (reads_df['trail_sc'] > 0)].shape[0]) if len(reads_df) else 0

# reads that reach the insertion locus and then clip toward the cassette
near_ins = 0
if len(reads_df) and 'ins0' in dir():
    tol = 6
    m5 = reads_df['clip5_boundary'].between(ins0 - tol, ins0 + tol)
    m3 = reads_df['clip3_boundary'].between(ins0 - tol, ins0 + tol)
    near_ins = int((m5 | m3).sum())

say('=' * 84)
say('VERDICT')
say('=' * 84)
say(f'  window aligned depth (max, any source): {max_depth_overall}')
say(f'  infected anchor reads: {inf_anchor}   negative anchor reads: {neg_anchor}')
say(f'  infected ground-truth core matches: {inf_gt}')
say(f'  infected reads soft-clipping (trailing): {inf_clip}')
say(f'  reads clipping AT the insertion locus (+/-6bp): {near_ins}')
say('')

if max_depth_overall == 0:
    say('  COVERAGE ZERO over the vpr window in every source.')
    say('  The reads are not there. Reference surgery cannot rescue reads that do')
    say('  not exist. This is the structural outcome: the barcode sits deep in the')
    say('  transcript interior, far from any 5-prime-captured end. Path forward is')
    say('  targeted VpxF1 cDNA enrichment (new sequencing), not reanalysis.')
elif near_ins > 0 and inf_gt > 0:
    say('  GREEN LIGHT for reference surgery. Reads reach the insertion locus and')
    say('  soft-clip toward the cassette, and at least one clipped/interior read')
    say('  canonicalizes to a real Binhua barcode. Adding the constant flanks +')
    say('  N10 core placeholder at the insertion locus should let these reads align')
    say('  through so the core can be read by coordinate. See Cell 7 for the plan.')
elif max_depth_overall > 0 and near_ins == 0 and inf_anchor == 0:
    say('  Reads cover the window but do NOT clip at the insertion locus and carry')
    say('  no anchor. Either coverage stops short of the cassette (5-prime bias, as')
    say('  suspected) or the anchor model is wrong. Inspect vpr_coverage.pdf: if the')
    say('  profile falls to zero before the insertion locus, it is the geometry and')
    say('  surgery will not help. If it reads through, recompute the anchor.')
else:
    say('  MIXED / borderline. There is some coverage and/or a few anchor reads but')
    say('  not a clean pileup-then-clip at the insertion locus. Read the two figures')
    say('  and vpr_window_reads.csv before deciding. Likely too sparse for reliable')
    say('  per-cell recovery even with surgery; quantify against the negative floor.')

say('')
say('  Specificity check: infected anchor reads should exceed the negative floor.')
say(f'    infected {inf_anchor} vs negative {neg_anchor}. If comparable, anchor hits')
say('    are background, not virus (matches the prior feasibility finding).')
say('=' * 84)

with open(os.path.join(OUT_DIR, 'verdict.txt'), 'w') as fh:
    fh.write('\n'.join(lines) + '\n')
print(f"\n  wrote {os.path.join(OUT_DIR, 'verdict.txt')}")


# %% ============================================================
# CELL 7 — surgery plan (documented, NOT executed here)
# ============================================================
hr('CELL 7 — reference surgery plan (only if Cell 6 is GREEN LIGHT)')
print("""
  If Cell 6 says GREEN LIGHT, the recovery reference is built like this (next
  session, separate script, not here):

   1. Take the reference 5-prime and 3-prime flanks printed in Cell 2 around the
      insertion locus (ins0). These are the real SHIVAD8EO sequence that reads
      align to on either side of the cassette.

   2. Build a cassette placeholder = [constant 5-prime flank from Binhua]
      + NNNNNNNNNN (the 10bp core, held as N so it never penalizes alignment)
      + [constant 3-prime flank from Binhua], using the consensus flanks Cell 3
      printed. Splice it into SHIVAD8EO at ins0 to make SHIVAD8EO_bc.fa.

   3. Rebuild the STARsolo index on SHIVAD8EO_bc.fa (star env, STAR 2.7.11b) and
      re-run the recovery alignment with the same relaxed filters. Reads that
      previously clipped the cassette now align through the constant flanks; the
      10bp core sits at a FIXED coordinate window we read straight off each read.

   4. Read the core by coordinate (not off the reference N's), canonicalize,
      cross-ref Binhua, and join CB -> animal/tissue/timepoint off the embedded
      object. That produces the barcode-by-cell matrix.

   The A3-hypermutation / pol-divergence problem is a SEPARATE, later lever: it
   affects how many SHIV reads map at all, upstream of the barcode locus. Note it
   as the next gap; do not fold it into the surgery.

  If Cell 6 was ZERO COVERAGE, skip all of the above and write the targeted-
  enrichment recommendation into the handoff instead.
""")
print('=' * 84)
print('vpr READ DIAGNOSTIC COMPLETE — read-only. No reference/object modified.')
print('=' * 84)
