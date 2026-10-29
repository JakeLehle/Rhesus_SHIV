#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Rhesus SHIV — vpr Cassette Flank Scan (Read-Only, Alignment-Free)
=================================================================
Finds the Keele viral barcode by anchoring on the VIRUS sequence that flanks the
cassette, not on the cassette's own half-sites, and not on the ground-truth
barcodes. Scans BOTH R1 cDNA and R2. Links every hit to a cell. Nothing on disk
is modified.

WHAT CHANGED FROM THE PRIOR PASSES, AND WHY
-------------------------------------------
shiv_vpr_barcode_demux.py anchored on 'ACGCGC.{22}GCGCGT', the cassette's own
MluI half-sites, and returned 3 reads. SHIV_vpr_barcode_gt_search.py added
partial and mismatch tiers against the Binhua set and added single digits.

Brandon Keele's advice (via Zoey, Oct 2026) was to stop using the barcode to
find the barcode and anchor on the flanking virus sequence instead. Three things
follow, and all three are real:

  1. ONE-SIDED ANCHORING. The old regex needed all 34bp of the cassette inside
     one read. A virus-flank anchor needs only the flank plus however many
     cassette bases the read happens to give. At 150bp that widens the window of
     usable read-start positions from 117 to 149.

  2. BARCODE-AGNOSTIC. The Binhua catalog is PLASMA-derived from specific draw
     dates. A cell-associated reservoir barcode need not be in it, so matching
     against it is a filter that can only lose reads. Here the catalog is a
     VALIDATION readout, never a gate.

  3. SPECIFICITY. 'ACGCGC' and 'GCGCGT' are 6-mers, expected roughly every 4kb
     by chance. A 20bp virus flank is essentially specific.

TWO ORIENTATION TRAPS, both of which would produce a believable zero
--------------------------------------------------------------------
  a) The old regex was its own reverse complement at the outer 6bp, so ONE
     forward pass caught both strands for free. Keele's flanks are NOT
     self-complementary, and shiv_starsolo_recovery.sh records that R2 is ~89%
     minus-strand. A forward-only search would lose nine reads in ten. Every
     probe here is searched in both orientations, explicitly.

  b) shiv_merge_barcode_diagnostic.py suspected a NovaSeq X cell-barcode flip.
     shiv_r1_structure_check.py has since measured the raw R1 CB as DIRECT, in
     whitelist orientation (~69% vs ~0.03%), so there is no FASTQ-level flip.
     That leaves the 458-reads-to-2-cells history unexplained, so CELL 6 still
     tests both orientations at join time and reports which wins.

WHY R1 IS SCANNED, AND WHY THAT IS NOT JUST A 2x
------------------------------------------------
STARsolo and CellRanger align R2 as the cDNA read and use R1 for nothing but CB
and UMI (--soloCBlen 16 --soloUMIlen 12; CB_UMI_Simple reads nothing past base
28). So the ~947-read SHIV pool this project has quoted throughout is an
R2-ONLY number, and R1 has never been aligned by anything here. A molecule whose
R2 failed to align to SHIV may still carry cassette sequence in its R1.

shiv_r1_structure_check.py measured R1 = 150bp, TSO at offset 28, cDNA from base
41, 109bp of it, poly-G 0.08-2.09%, entropy flat at ~1.90 bits across all five
20bp blocks. That is real transcript, not dark-cycle artifact.

This script scans R1[28:] ENTIRE (122bp) rather than from the measured base 41,
because 13-21% of reads had no TSO detectable at any offset yet still carry a
valid CB and usable sequence. The TSO is a constant 13-mer and cannot match a
virus flank, so including it costs nothing and removes a per-read dependency on
TSO detection. Searchable sequence per read pair is therefore 122 + 150 = 272bp.

AND THE PRIZE: a cassette hit in R1 carries the cell barcode in the SAME
molecule, 28 bases upstream. No pairing assumption, no orientation decision, no
join. That is the strongest barcode-to-cell evidence this dataset can produce,
and it is given its own evidence tier above everything else.

HONEST EXPECTATION
------------------
The R2 arm is predicted at ~13.8 reads dataset-wide (shiv_read_length_check.py
CELL 4, at the measured 150bp). The R1 arm cannot be predicted from that model,
because R1 always samples the molecule's 5' end rather than a uniform position,
so it has to be measured. Do not run this expecting a reservoir map. Run it
because it is the only measurement in this project that is independent of every
alignment decision made so far, which is what makes the enrichment case
defensible if it comes back empty.

Conventions: Spyder cells (# %%), config in Cell 1, warn-and-continue,
read-only. Runs headless under SLURM.

Author: Jake Lehle
Date: October 2026
"""

# %% ============================================================
# CELL 1 — CONFIG AND DIALS
# ============================================================
import os
import re
import glob
import shutil
import subprocess
from collections import Counter, defaultdict

import pandas as pd

# --- paths -----------------------------------------------------------------
WORKING   = '/master/jlehle/WORKING/SC/fastq/Rhesus_SHIV'
RAW_BASE  = '/master/zwallis/WORKING/SC/10X_RAW'
LIB_DIR   = 'GEX_LIBRARY_{L}'
OBJECT    = os.path.join(WORKING, 'shiv_starsolo', 'merged',
                         'shiv_host_final_embedded.h5ad')
BINHUA    = os.path.join(WORKING,
                         'Binhua_FINAL_Barcode_4-7-26_Rhesus_SHIV_Barcodes.xlsx')
BINHUA_ALT = os.path.join(WORKING, 'Binhua_FINAL_Barcode_4726.xlsx')
OUT_DIR   = os.path.join(WORKING, 'annotation_output', 'vpr_flank_scan')

# --- libraries -------------------------------------------------------------
LIB_NAME = {
    'A': 'SHIV_Pre_PBMC',             'B': 'SHIV_Pre_LN',
    'C': 'SHIV_Pre_PBMC_IndianMixed', 'D': 'SHIV_21DPI_Pre_LN_Mixed',
    'E': 'SHIV_21DPI_PBMC',           'F': 'SHIV_21DPI_NonInf_LN_Mixed',
    'G': 'SHIV_Necropsy_PBMC',        'H': 'SHIV_Necropsy_LN',
}
LIBS     = ['A', 'B', 'C', 'D', 'E', 'F', 'G', 'H']   # negatives ARE the floor
INFECTED = {'D', 'E', 'F', 'H'}

# --- SMOKE TEST FIRST ------------------------------------------------------
# Run once with MAX_READS_PER_LIB = 10_000_000 and LIBS = ['E','H'] to confirm
# the probe build, the decompression pipe and the CB join all behave. Then set
# to None for the full pass. Do not discover a typo at hour six.
MAX_READS_PER_LIB = None          # None = full scan; int = cap per library
PROGRESS_EVERY    = 25_000_000

# --- Brandon Keele's flanking virus sequence (Oct 2026) --------------------
# Verified against his SHIVAD8EOM reference: the 5' flank ends exactly at the
# 34-N run and the 3' flank begins exactly at its end, no off-by-one. In OUR
# reference frame (MN816822.1) the insertion locus is base 5862.
KEELE_FLANK_5 = 'GGAGACCAGGACCTCCTCCTCCTCCCCCTCCAGGACTAGCATAA'   # 44bp
KEELE_FLANK_3 = 'ATGGAAGAAAGACCTCCAGAAAATGAAGGACCACAAAGGGAA'     # 42bp
INSERTION_LOCUS_OURS = 5862       # 0-based, for cross-referencing the BAM work

MIN_VIRUS_FLANK = 20              # bp of virus flank required to anchor
SEED_LEN        = 12              # cheap byte prefilter, cassette-proximal end

# --- cassette geometry (Khanal et al. 2019, J Virol; confirmed against the
#     Binhua workbook by shiv_vpr_barcode_feasibility.py) -------------------
BC_LEN        = 34                # 12bp flank + 10bp core + 12bp flank
CFLANK        = 12
CORE_LEN      = 10
INNER_ORIENT1 = 'AGCGTA'          # core taken as-is
INNER_ORIENT2 = 'TAGCTG'          # core reverse-complemented to canonicalize
MIN_CORE_BASES = 6                # of the 10bp core, how many must be visible

# --- 10x GEM-X 5' v3 read geometry (measured, not assumed) ------------------
#     shiv_read_length_check.py : R1 150bp, R2 150bp, uniform, 8/8 libraries
#     shiv_r1_structure_check.py: TSO offset 28, cDNA from 41, CB direct
CB_LEN        = 16
UMI_LEN       = 12
R1_SCAN_FROM  = CB_LEN + UMI_LEN  # 28; scan R1[28:] entire, TSO included
SCAN_R1       = True              # set False to reproduce the R2-only passes
SCAN_R2       = True

# --- poly-G guard (NovaSeq X dark-cycle artifact) --------------------------
POLYG_RUN = 15                    # trim R1 cDNA at the first G-run this long

# --- ground truth ----------------------------------------------------------
USE_GT_AS_VALIDATION_ONLY = True  # True = never gate on the catalog. Keep True.

os.makedirs(OUT_DIR, exist_ok=True)
pd.set_option('display.width', 180)
pd.set_option('display.max_columns', 60)

_COMP = str.maketrans('ACGTNacgtn', 'TGCANtgcan')


def revcomp(s):
    return s.translate(_COMP)[::-1]


def hr(title):
    print("\n" + "=" * 86)
    print(title)
    print("=" * 86)


def save_csv(df, name, index=False):
    p = os.path.join(OUT_DIR, name)
    df.to_csv(p, index=index)
    print(f"    saved: {name}  ({len(df)} rows)")


def open_fastq(path):
    """Stream a .fastq.gz through pigz or zcat. Python's gzip module is the
    bottleneck at this scale (~130GB compressed across 16 files), so hand the
    decompression to a C implementation and read its stdout."""
    for tool, args in (('pigz', ['-dc']), ('zcat', [])):
        exe = shutil.which(tool)
        if exe:
            p = subprocess.Popen([exe] + args + [path],
                                 stdout=subprocess.PIPE,
                                 stderr=subprocess.DEVNULL,
                                 bufsize=1024 * 1024)
            return p, p.stdout
    import gzip                                   # last resort
    print(f"    WARN  neither pigz nor zcat found; falling back to "
          f"python gzip (slow) for {os.path.basename(path)}")
    return None, gzip.open(path, 'rb')


# %% ============================================================
# CELL 2 — BUILD THE PROBE SET
# ============================================================
hr("CELL 2 — probe construction")

# The cassette-proximal end of each virus flank is what must abut the barcode:
# for the 5' flank that is its 3' terminus, for the 3' flank its 5' terminus.
PROBE_5 = KEELE_FLANK_5[-MIN_VIRUS_FLANK:]
PROBE_3 = KEELE_FLANK_3[:MIN_VIRUS_FLANK]
SEED_5  = KEELE_FLANK_5[-SEED_LEN:]
SEED_3  = KEELE_FLANK_3[:SEED_LEN]

print(f"  Keele 5' flank      : {KEELE_FLANK_5}  ({len(KEELE_FLANK_5)}bp)")
print(f"  Keele 3' flank      : {KEELE_FLANK_3}  ({len(KEELE_FLANK_3)}bp)")
print(f"  insertion locus     : {INSERTION_LOCUS_OURS} (0-based, our frame)")
print(f"\n  probe 5' (last {MIN_VIRUS_FLANK}bp) : {PROBE_5}")
print(f"  probe 3' (first {MIN_VIRUS_FLANK}bp): {PROBE_3}")
print(f"  seed 5'             : {SEED_5}   revcomp {revcomp(SEED_5)}")
print(f"  seed 3'             : {SEED_3}   revcomp {revcomp(SEED_3)}")

# Byte-level prefilter. One compiled alternation over the four seeds; on a hit
# we do the full work. Over ~800M read pairs this gate is the whole runtime.
SEEDS_B = [s.encode() for s in
           (SEED_5, SEED_3, revcomp(SEED_5), revcomp(SEED_3))]
SEED_RE_B = re.compile(b'|'.join(re.escape(s) for s in SEEDS_B))
REV_SEEDS_B = {revcomp(SEED_5).encode(), revcomp(SEED_3).encode()}

print(f"\n  prefilter: {len(SEEDS_B)} seeds, both orientations, "
      f"compiled alternation")
print(f"  scanning  : R1[{R1_SCAN_FROM}:] = {SCAN_R1}   R2 = {SCAN_R2}")

# Sanity: the probes must not be self-complementary, or the two-orientation
# logic below collapses and we would silently double-count.
for nm, p in (('5prime', PROBE_5), ('3prime', PROBE_3)):
    if p == revcomp(p):
        print(f"  *** WARN {nm} probe is its own reverse complement; "
              f"orientation calls will be ambiguous")


def canon_core(bc):
    """Given forward-strand cassette bases, return (canonical_core, n_core_bases,
    full) where full is True only for a complete, well-formed 34bp cassette.

    Register is read off the inner flank: AGCGTA means the core is as-is,
    TAGCTG means revcomp it. Mirrors canon_core() in
    shiv_vpr_barcode_feasibility.py so recovered cores are comparable to every
    prior pass, but tolerates a partial cassette instead of returning None.
    """
    if len(bc) < CFLANK + MIN_CORE_BASES:
        return None, max(0, len(bc) - CFLANK), False
    inner = bc[CFLANK - 6:CFLANK]
    core = bc[CFLANK:CFLANK + CORE_LEN]
    n = len(core)
    if inner == INNER_ORIENT1:
        pass
    elif inner == INNER_ORIENT2:
        core = revcomp(core)
    else:
        return None, n, False                     # corrupted register
    return core, n, (len(bc) >= BC_LEN and n == CORE_LEN)


# %% ============================================================
# CELL 3 — GROUND TRUTH (VALIDATION ONLY, NEVER A GATE)
# ============================================================
hr("CELL 3 — Binhua ground truth (validation readout, not a filter)")

gt_core_to_animals = defaultdict(set)
gt_path = BINHUA if os.path.exists(BINHUA) else BINHUA_ALT
if not os.path.exists(gt_path):
    print(f"  WARN  no Binhua workbook found at either path.")
    print("        The scan still runs; GT columns will be empty.")
else:
    try:
        from openpyxl import load_workbook
        wb = load_workbook(gt_path, read_only=True, data_only=True)
        n_rows = n_ok = 0
        for sh in wb.sheetnames:
            animal = sh.split()[0]
            for row in wb[sh].iter_rows(values_only=True):
                v = row[2] if len(row) > 2 else None
                if not isinstance(v, str):
                    continue
                s = v.strip().upper()
                if not re.fullmatch(r'[ACGTN]{20,40}', s):
                    continue
                n_rows += 1
                core, _, full = canon_core(s)
                if full and core:
                    gt_core_to_animals[core].add(animal)
                    n_ok += 1
        print(f"  workbook        : {os.path.basename(gt_path)}")
        print(f"  rows parsed     : {n_rows}")
        print(f"  canonical cores : {n_ok} -> {len(gt_core_to_animals)} unique")
        shared = {c: a for c, a in gt_core_to_animals.items() if len(a) > 1}
        print(f"  cores in >1 animal: {len(shared)}")
        print("    Animal concordance in CELL 7 is computed on animal-PRIVATE "
              "cores only;\n    a core present in several independently "
              "inoculated animals cannot\n    attribute a cell to one of them.")
    except Exception as e:                        # noqa: BLE001
        print(f"  WARN  could not parse workbook: {type(e).__name__}: {e}")

GT_PRIVATE = {c: next(iter(a)) for c, a in gt_core_to_animals.items()
              if len(a) == 1}


# %% ============================================================
# CELL 4 — THE SCAN
# ============================================================
hr("CELL 4 — streaming scan of raw R1 + R2")

TIERS = ['R1_FULL', 'R1_PARTIAL', 'R2_FULL', 'R2_PARTIAL', 'FLANK_ONLY']


def trim_polyg(s):
    """Cut R1 cDNA at the first long G-run. NovaSeq X calls G on dark cycles,
    so a poly-G tail is absence of signal, not sequence."""
    i = s.find('G' * POLYG_RUN)
    return s[:i] if i >= 0 else s


def extract(seq, which_read):
    """Find a virus flank in seq (either orientation) and return the cassette
    bases adjacent to it. Reverse hits are handled by reverse-complementing the
    whole read and re-finding, which is far less error-prone than mirroring the
    index arithmetic. Returns a dict or None."""
    best = None
    for orient, s in (('fwd', seq), ('rev', revcomp(seq))):
        i5 = s.find(PROBE_5)
        i3 = s.find(PROBE_3)
        if i5 < 0 and i3 < 0:
            continue
        # Prefer whichever side yields more cassette bases.
        cands = []
        if i5 >= 0:
            cands.append(('5prime', s[i5 + len(PROBE_5):][:BC_LEN]))
        if i3 >= 0:
            tail = s[:i3]
            cands.append(('3prime', tail[-BC_LEN:] if tail else ''))
        for side, bc in cands:
            core, n_core, full = canon_core(bc)
            both = (i5 >= 0 and i3 >= 0)
            cand = {'read': which_read, 'orient': orient, 'side': side,
                    'both_flanks': both, 'cassette_bases': len(bc),
                    'cassette': bc, 'core': core, 'n_core': n_core,
                    'full': full}
            if best is None or (cand['n_core'], cand['full']) > \
                               (best['n_core'], best['full']):
                best = cand
    return best


def tier_of(hit):
    if hit['n_core'] < MIN_CORE_BASES or hit['core'] is None:
        return 'FLANK_ONLY'
    if hit['read'] == 'R1':
        return 'R1_FULL' if hit['full'] else 'R1_PARTIAL'
    return 'R2_FULL' if hit['full'] else 'R2_PARTIAL'


hits = []
per_lib = []

for L in LIBS:
    d = os.path.join(RAW_BASE, LIB_DIR.format(L=L))
    r1s = sorted(glob.glob(os.path.join(d, '*_R1_*.fastq.gz')))
    r2s = sorted(glob.glob(os.path.join(d, '*_R2_*.fastq.gz')))
    if not r1s or not r2s:
        print(f"  {L}: missing R1/R2 in {d}; skipped")
        continue
    if len(r1s) != len(r2s):
        print(f"  {L}: WARN R1/R2 file-count mismatch "
              f"({len(r1s)}/{len(r2s)}); pairing by sort order")

    n_reads = n_seed = 0
    tier_counts = Counter()
    grp = 'INF' if L in INFECTED else 'neg'

    for p1, p2 in zip(r1s, r2s):
        proc1, f1 = open_fastq(p1)
        proc2, f2 = open_fastq(p2)
        try:
            while True:
                h1 = f1.readline()
                if not h1:
                    break
                s1 = f1.readline().rstrip(b'\n')
                f1.readline(); f1.readline()
                f2.readline()
                s2 = f2.readline().rstrip(b'\n')
                f2.readline(); f2.readline()
                n_reads += 1

                # --- the gate: one regex over both reads' searchable spans ---
                span1 = s1[R1_SCAN_FROM:] if SCAN_R1 else b''
                span2 = s2 if SCAN_R2 else b''
                if not (SEED_RE_B.search(span1) or SEED_RE_B.search(span2)):
                    if MAX_READS_PER_LIB and n_reads >= MAX_READS_PER_LIB:
                        break
                    if n_reads % PROGRESS_EVERY == 0:
                        print(f"    {L}: {n_reads:,} reads, "
                              f"{n_seed} seed hits")
                    continue
                n_seed += 1

                cb = s1[:CB_LEN].decode(errors='ignore')
                umi = s1[CB_LEN:CB_LEN + UMI_LEN].decode(errors='ignore')
                rid = h1.split()[0].decode(errors='ignore').lstrip('@')

                found = []
                if SCAN_R1 and SEED_RE_B.search(span1):
                    cdna = trim_polyg(span1.decode(errors='ignore'))
                    h = extract(cdna, 'R1')
                    if h:
                        found.append(h)
                if SCAN_R2 and SEED_RE_B.search(span2):
                    h = extract(s2.decode(errors='ignore'), 'R2')
                    if h:
                        found.append(h)

                for h in found:
                    t = tier_of(h)
                    tier_counts[t] += 1
                    core = h['core']
                    hits.append({
                        'lib': L, 'group': grp, 'read_id': rid,
                        'cb': cb, 'umi': umi,
                        'tier': t, **{k: v for k, v in h.items()
                                      if k != 'core'},
                        'core': core,
                        'gt_animal': GT_PRIVATE.get(core) if core else None,
                        'gt_known': (core in gt_core_to_animals) if core
                                    else False,
                    })

                if MAX_READS_PER_LIB and n_reads >= MAX_READS_PER_LIB:
                    break
                if n_reads % PROGRESS_EVERY == 0:
                    print(f"    {L}: {n_reads:,} reads, {n_seed} seed hits")
        finally:
            for pr, fh in ((proc1, f1), (proc2, f2)):
                try:
                    fh.close()
                except Exception:                 # noqa: BLE001
                    pass
                if pr is not None:
                    pr.wait()
        if MAX_READS_PER_LIB and n_reads >= MAX_READS_PER_LIB:
            break

    row = {'lib': L, 'group': grp, 'sample': LIB_NAME[L],
           'n_reads': n_reads, 'n_seed_hits': n_seed}
    row.update({t: tier_counts.get(t, 0) for t in TIERS})
    row['n_called'] = sum(tier_counts.get(t, 0) for t in TIERS
                          if t != 'FLANK_ONLY')
    per_lib.append(row)
    print(f"  {L} [{grp}]: {n_reads:,} reads, {n_seed} seed, "
          f"called {row['n_called']}, flank-only "
          f"{tier_counts.get('FLANK_ONLY', 0)}")

hit_df = pd.DataFrame(hits)
lib_df = pd.DataFrame(per_lib)
if not lib_df.empty:
    save_csv(lib_df, 'flank_scan_per_library.csv')
    print("\n" + lib_df.to_string(index=False))
if not hit_df.empty:
    save_csv(hit_df, 'flank_scan_hits.csv')


# %% ============================================================
# CELL 5 — NEGATIVE FLOOR
# ============================================================
hr("CELL 5 — specificity against the negative floor")

if lib_df.empty:
    print("  nothing scanned")
else:
    g = (lib_df.groupby('group')
         .agg(n_reads=('n_reads', 'sum'), n_called=('n_called', 'sum'),
              flank_only=('FLANK_ONLY', 'sum'),
              **{t: (t, 'sum') for t in TIERS if t != 'FLANK_ONLY'})
         .reset_index())
    for c in ['n_called', 'flank_only']:
        g[c + '_per_1e9'] = (g[c] / g['n_reads'] * 1e9).round(2)
    print(g.to_string(index=False))
    save_csv(g, 'flank_scan_by_group.csv')

    inf = g[g['group'] == 'INF']
    neg = g[g['group'] == 'neg']
    if not inf.empty and not neg.empty:
        i = inf['n_called_per_1e9'].iloc[0]
        n = neg['n_called_per_1e9'].iloc[0]
        print(f"\n  called rate per 1e9 reads: infected {i}, negative {n}")
        if n == 0 and i > 0:
            print("  Clean. Negatives are empty, so every infected call is "
                  "above background.")
        elif n > 0 and i / max(n, 1e-9) < 3:
            print("  *** Infected rate is within 3x of the negative floor. "
                  "These are\n      background, not virus. Same conclusion as "
                  "the prior feasibility pass.")
        else:
            print(f"  Infected exceeds the floor by "
                  f"{i / max(n, 1e-9):.1f}x.")

    print("\n  FLANK_ONLY is the new diagnostic. Reads carrying >=20bp of virus"
          "\n  flank but too few cassette bases to call. A high count means "
          "reads DO\n  reach the locus and simply do not extend into the "
          "cassette; near-zero\n  means the locus is genuinely unsampled, "
          "which closes the question.")


# %% ============================================================
# CELL 6 — CELL LINKAGE (both CB orientations, empirically)
# ============================================================
hr("CELL 6 — join to called cells")

obs = None
if hit_df.empty:
    print("  no hits to link")
elif not os.path.exists(OBJECT):
    print(f"  WARN  object not found at {OBJECT}; linkage skipped")
else:
    try:
        import anndata as ad
        adata = ad.read_h5ad(OBJECT)
        obs = adata.obs
        print(f"  object: {adata.shape[0]:,} cells x {adata.shape[1]:,} genes")

        # GEX obs_names are '<bare>-1-<lib>' (shiv_VDJ_integration.py).
        bare_by_lib = defaultdict(set)
        for name in adata.obs_names:
            parts = name.split('-')
            if len(parts) >= 3:
                bare_by_lib[parts[-1]].add(parts[0])
            else:
                bare_by_lib['?'].add(parts[0])
        print(f"  cell barcodes indexed for libraries: "
              f"{sorted(bare_by_lib)}")

        # Test both orientations rather than trusting either. The r1 structure
        # check says raw CB is direct, but the 458-to-2 merge history is
        # unexplained, so this stays empirical.
        res = {}
        for mode in ('direct', 'revcomp'):
            n = 0
            for _, r in hit_df.iterrows():
                cb = r['cb'] if mode == 'direct' else revcomp(r['cb'])
                if cb in bare_by_lib.get(r['lib'], ()):
                    n += 1
            res[mode] = n
        print(f"\n  CB join yield: direct {res['direct']}/{len(hit_df)}, "
              f"revcomp {res['revcomp']}/{len(hit_df)}")
        MODE = max(res, key=res.get)
        if res[MODE] == 0:
            print("  *** Neither orientation joins. Do NOT interpret the hits "
                  "as cells.")
        else:
            print(f"  using {MODE}")

        keep = []
        for _, r in hit_df.iterrows():
            cb = r['cb'] if MODE == 'direct' else revcomp(r['cb'])
            key = f"{cb}-1-{r['lib']}"
            if key in obs.index:
                o = obs.loc[key]
                keep.append({
                    **r.to_dict(), 'cell': key,
                    'animal_id': o.get('animal_id', 'Unknown'),
                    'tissue': o.get('tissue', None),
                    'timepoint': o.get('timepoint', None),
                    'tier1_celltype': o.get('tier1_celltype', None),
                    'shiv_pos': o.get('shiv_pos', None),
                })
        linked = pd.DataFrame(keep)
        if linked.empty:
            print("  no hit joined a called cell")
        else:
            save_csv(linked, 'flank_scan_linked_cells.csv')
            print(f"\n  linked {len(linked)} hit(s) to "
                  f"{linked['cell'].nunique()} cell(s)")
            print(linked[['lib', 'group', 'tier', 'core', 'animal_id',
                          'tissue', 'timepoint', 'tier1_celltype']]
                  .to_string(index=False))
    except Exception as e:                        # noqa: BLE001
        print(f"  WARN  linkage failed: {type(e).__name__}: {e}")
        linked = pd.DataFrame()


# %% ============================================================
# CELL 7 — VALIDATION AND VERDICT
# ============================================================
hr("CELL 7 — validation and verdict")

if hit_df.empty:
    print("  No hits anywhere.")
    print("\n  With R1 and R2 both scanned, both orientations covered, no")
    print("  ground-truth gate and no dependence on any alignment, this is the")
    print("  most permissive probe the data admits. A zero here is a direct")
    print("  measurement that the cassette locus is unsampled in this library")
    print("  set, not an inference from aligned counts. That is the result")
    print("  that makes the cDNA-based VpxF1 enrichment case, and it should go")
    print("  in the methods as a stated negative.")
else:
    print("  hits by tier:")
    for t in TIERS:
        sub = hit_df[hit_df['tier'] == t]
        if len(sub):
            print(f"    {t:12s} {len(sub):5d}   "
                  f"INF {len(sub[sub['group'] == 'INF']):4d}  "
                  f"neg {len(sub[sub['group'] == 'neg']):4d}")

    called = hit_df[hit_df['tier'] != 'FLANK_ONLY']
    if len(called):
        gt_rate = called['gt_known'].mean()
        print(f"\n  cores matching the Binhua catalog: "
              f"{called['gt_known'].sum()}/{len(called)} ({gt_rate:.0%})")
        print("    Not a pass/fail. The catalog is plasma-derived, so a "
              "reservoir core\n    absent from it is expected rather than "
              "suspect. High agreement is\n    reassurance; low agreement with "
              "a clean negative floor is a finding.")

    r1 = hit_df[hit_df['tier'].str.startswith('R1')]
    if len(r1):
        print(f"\n  *** {len(r1)} hit(s) in R1. These carry the cell barcode "
              f"in the SAME\n      molecule {R1_SCAN_FROM} bases upstream, so "
              "the cell assignment needs no\n      pairing assumption and no "
              "orientation decision. Treat these as the\n      highest-"
              "confidence evidence in the set.")

    try:
        if not linked.empty and 'animal_id' in linked:
            priv = linked[linked['gt_animal'].notna()]
            if len(priv):
                conc = (priv['gt_animal'].astype(str) ==
                        priv['animal_id'].astype(str)).mean()
                print(f"\n  animal concordance on animal-private cores: "
                      f"{conc:.0%} of {len(priv)}")
                print("    Barcode-implied animal vs the object's genetic-demux "
                      "animal_id. These\n    are independent, so agreement "
                      "supports both; disagreement says which\n    one to "
                      "distrust. Libraries C, D and F are mixed pools, so "
                      "animal_id\n    never comes from the library.")
    except NameError:
        pass

hr("FLANK SCAN COMPLETE — read-only. No reference, BAM or object modified.")
